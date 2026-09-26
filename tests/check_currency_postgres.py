import argparse
from decimal import Decimal
from pathlib import Path
import socket
import subprocess
import tempfile
from unittest.mock import patch

import psycopg2

from check_currency_workflow import make_test_app, check_api


def check_migration(application, connection):
    with application.app.app_context():
        product = application.Product(name='Legacy migration check', cost_price=4, sell_price=8)
        sale = application.Sale(customer_id=1, total_amount=66, debt_usd=66,
                                debt_amount=825000, currency_rate=12500, payment_status='debt')
        application.db.session.add_all([product, sale])
        application.db.session.commit()
        product_id, sale_id = product.id, sale.id
        application.db.session.remove()
    columns = {
        'products': ['cost_currency', 'sell_currency', 'native_cost_price', 'native_sell_price',
                     'native_last_batch_cost', 'price_currency_rate'],
        'sale_items': ['price_currency', 'native_unit_price', 'cost_currency', 'native_cost_price'],
        'sales': ['native_debt_usd', 'native_debt_uzs', 'native_payments'],
        'debt_payments': ['native_allocation'],
        'product_add_history': ['cost_currency', 'sell_currency', 'native_cost_price',
                                'native_sell_price', 'currency_rate']}
    from psycopg2 import sql
    with connection.cursor() as cursor:
        for table, names in columns.items():
            for name in names:
                cursor.execute(sql.SQL('ALTER TABLE {} DROP COLUMN {}').format(
                    sql.Identifier(table), sql.Identifier(name)))
        cursor.execute('ALTER TABLE products ALTER COLUMN cost_price TYPE NUMERIC(20,4)')
        cursor.execute('ALTER TABLE sales ALTER COLUMN debt_usd TYPE NUMERIC(20,4)')
        migration = (Path(__file__).resolve().parents[1] / 'migrations/add_native_currencies.sql').read_text()
        cursor.execute(migration)
        cursor.execute(migration)
        cursor.execute('SELECT cost_price, sell_price, native_cost_price, cost_currency FROM products WHERE id=%s', (product_id,))
        assert cursor.fetchone() == (Decimal('4'), Decimal('8'), None, 'USD')
        cursor.execute('SELECT debt_usd, debt_amount, native_debt_usd, native_debt_uzs FROM sales WHERE id=%s', (sale_id,))
        assert cursor.fetchone() == (Decimal('66'), Decimal('825000'), None, None)
        cursor.execute("SELECT numeric_precision, numeric_scale FROM information_schema.columns WHERE table_name='sales' AND column_name='debt_usd'")
        assert cursor.fetchone() == (24, 10)
        cursor.execute('DELETE FROM sales WHERE id=%s', (sale_id,))
        cursor.execute('DELETE FROM products WHERE id=%s', (product_id,))
    print('PostgreSQL migration: repeatable, legacy amounts unchanged, numeric precision OK')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--bin', default=r'C:\Program Files\PostgreSQL\17\bin')
    options = parser.parse_args()
    binaries = Path(options.bin)
    if not (binaries / 'initdb.exe').is_file():
        parser.error('PostgreSQL initdb.exe not found; supply --bin')
    with socket.socket() as listener:
        listener.bind(('127.0.0.1', 0))
        port = listener.getsockname()[1]
    with tempfile.TemporaryDirectory(prefix='currency-postgres-') as directory:
        root = Path(directory)
        cluster = root / 'data'
        subprocess.run([str(binaries / 'initdb.exe'), '-D', str(cluster), '-U', 'currency_test',
                        '-A', 'trust', '--encoding=UTF8', '--no-locale'], check=True, capture_output=True)
        started = False
        application = None
        connection = None
        try:
            started = True
            subprocess.run([str(binaries / 'pg_ctl.exe'), '-D', str(cluster), '-l', str(root / 'server.log'),
                            '-o', f'-h 127.0.0.1 -p {port}', '-w', 'start'], check=True,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=45)
            config = {'host': '127.0.0.1', 'port': port, 'user': 'currency_test', 'dbname': 'postgres'}
            connection = psycopg2.connect(**config)
            connection.autocommit = True
            application = make_test_app(f'postgresql://currency_test@127.0.0.1:{port}/postgres')
            check_migration(application, connection)
            check_api(application)
            response = application.app.test_client().get('/api/debts')
            assert response.status_code == 200, response.get_json()
            print('PostgreSQL real API and debts list OK')
            import export_excel_backup
            import openpyxl
            output = root / 'currency-report.xlsx'
            with patch.multiple(export_excel_backup, DB_CONFIG=config,
                                OUTPUT_DIR=str(root), OUTPUT=str(output)):
                export_excel_backup.main()
            workbook = openpyxl.load_workbook(output, data_only=True)
            debt_sheet = workbook.worksheets[2]
            assert debt_sheet.cell(2, 4).value == 50
            assert debt_sheet.cell(2, 5).value == 100000
            product_sheet = workbook.worksheets[0]
            som_row = next(row for row in product_sheet.iter_rows(values_only=True) if row[1] == 'UZS product')
            assert som_row[10:12] == (100000, 'UZS')
            workbook.close()
            print('PostgreSQL Excel: original product prices and independent debts OK')
        finally:
            if connection:
                connection.close()
            if application:
                with application.app.app_context():
                    application.db.session.remove()
                    application.db.engine.dispose()
            if started:
                subprocess.run([str(binaries / 'pg_ctl.exe'), '-D', str(cluster), '-m', 'fast', '-w', 'stop'],
                               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=45)


if __name__ == '__main__':
    main()
